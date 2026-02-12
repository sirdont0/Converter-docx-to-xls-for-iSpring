import os
from docx import Document
from openpyxl import Workbook
import tkinter as tk
from tkinter import filedialog, messagebox


COLUMNS = [
    "Тип вопроса",
    "Текст вопроса",
    "Рисунок",
    "Видео",
    "Аудио",
    "Ответ 1",
    "Ответ 2",
    "Ответ 3",
    "Ответ 4",
    "Ответ 5",
    "Ответ 6",
    "Ответ 7",
    "Ответ 8",
    "Ответ 9",
    "Ответ 10",
    "Сообщение, если верно",
    "Сообщение, если неверно",
    "Баллы"
]


def parse_table_from_docx(file_path):
    doc = Document(file_path)

    questions = []

    for table in doc.tables:
        for row in table.rows[1:]:  # пропускаем заголовок
            cells = [cell.text.strip() for cell in row.cells]

            if len(cells) < 5:
                continue

            formulation = cells[2]
            variants = cells[3]
            correct = cells[4]

            if not formulation:
                continue

            answers = [""] * 10

            # ----- MC -----
            if variants:
                q_type = "MC"

                lines = [v.strip() for v in variants.split("\n") if v.strip()]
                letter_map = {}

                for idx, line in enumerate(lines):
                    if ")" in line:
                        letter = line[0]
                        text = line.split(")", 1)[1].strip()
                        letter_map[letter] = idx
                        answers[idx] = text

                # ставим звёздочку ПЕРЕД ответом без пробела
                correct_letters = [c.strip() for c in correct.split(",")]

                for letter in correct_letters:
                    if letter in letter_map:
                        idx = letter_map[letter]
                        answers[idx] = "*" + answers[idx]

            # ----- TI -----
            else:
                q_type = "TI"

                # заменяем все разделители на перенос строки
                clean_text = correct.replace(",", "\n").replace(";", "\n")

                # разбиваем и очищаем
                ti_answers = [a.strip() for a in clean_text.split("\n") if a.strip()]

                for idx, ans in enumerate(ti_answers):
                    if idx < 10:
                        answers[idx] = ans

            row_data = [
                q_type,
                formulation,
                "", "", "",
                *answers,
                "", "",
                1
            ]

            questions.append(row_data)

    return questions


def create_excel(file_path, questions):
    wb = Workbook()
    ws = wb.active

    ws.append(COLUMNS)

    for q in questions:
        ws.append(q)

    output_file = os.path.splitext(file_path)[0] + "_iSpring.xlsx"
    wb.save(output_file)


def select_files():
    root = tk.Tk()
    root.withdraw()

    files = filedialog.askopenfilenames(
        title="Выберите Word-файлы",
        filetypes=[("Word files", "*.docx")]
    )

    if not files:
        return

    for file in files:
        questions = parse_table_from_docx(file)
        create_excel(file, questions)

    messagebox.showinfo("Готово", "Файлы успешно конвертированы!")


if __name__ == "__main__":
    select_files()