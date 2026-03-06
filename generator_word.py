import os
import json
from datetime import datetime
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import tkinter as tk
from tkinter import filedialog

DEBUG = True
TEMPLATE_PATH = "./Шаблон_теста.docx"


def debug_log(stage, message):
    if not DEBUG:
        return
    timestamp = datetime.now().strftime("%H:%M:%S")
    print(f"[{timestamp}] [{stage}] {message}", flush=True)


def clean_json_text(raw_text):
    text = raw_text.strip()
    text = text.replace("```json", "").replace("```", "").strip()
    text = text.replace("“", '"').replace("”", '"')

    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end > start:
        text = text[start:end + 1]
    return text


def load_questions_from_json(json_path):
    debug_log("JSON", f"Чтение файла: {json_path}")
    with open(json_path, "r", encoding="utf-8-sig") as source:
        raw_text = source.read()

    cleaned = clean_json_text(raw_text)
    data = json.loads(cleaned)

    if "closed_questions" not in data or "open_questions" not in data:
        raise ValueError("JSON должен содержать поля 'closed_questions' и 'open_questions'")

    debug_log(
        "JSON",
        f"Данные загружены: closed={len(data.get('closed_questions', []))}, "
        f"open={len(data.get('open_questions', []))}"
    )
    return data


def center_cell_content(cell):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def get_document_title(data, fallback_name):
    title = data.get("tittle_name") or data.get("title_name") or fallback_name
    title = str(title).strip()
    return title if title else fallback_name


def sanitize_filename(name):
    invalid_chars = '<>:"/\\|?*'
    sanitized = "".join("_" if ch in invalid_chars else ch for ch in name).strip().rstrip(".")
    return sanitized if sanitized else "test"


def fill_word_template(template_path, output_path, document_name, data):

    debug_log("WORD", f"Открытие шаблона: {template_path}")
    doc = Document(template_path)

    # ======================
    # вставляем название документа
    # ======================

    debug_log("WORD", f"Вставка названия документа: {document_name}")
    for paragraph in doc.paragraphs:
        if "1 – 16" in paragraph.text:
            paragraph.text = f"1 – 16 {document_name}"
            break

    # ======================
    # таблица
    # ======================

    debug_log("WORD", "Заполнение таблицы вопросами...")
    table = doc.tables[0]

    row_index = 1
    question_number = 1

    # закрытые
    debug_log("WORD", f"Закрытых вопросов: {len(data['closed_questions'])}")
    for q in data["closed_questions"]:

        row = table.rows[row_index]

        row.cells[0].text = "1" if question_number == 1 else ""
        row.cells[1].text = f"{question_number}."
        row.cells[2].text = q["question"]

        variants = ""
        for key, value in q["options"].items():
            variants += f"{key}) {value}\n"

        row.cells[3].text = variants.strip()
        row.cells[4].text = q["correct"]
        center_cell_content(row.cells[4])

        row_index += 1
        question_number += 1

    # открытые
    debug_log("WORD", f"Открытых вопросов: {len(data['open_questions'])}")
    for q in data["open_questions"]:

        row = table.rows[row_index]

        row.cells[0].text = "1" if question_number == 1 else ""
        row.cells[1].text = f"{question_number}."
        row.cells[2].text = q["question"]
        row.cells[3].text = ""
        row.cells[4].text = q["answer"]
        center_cell_content(row.cells[4])

        row_index += 1
        question_number += 1

    debug_log("WORD", f"Сохранение результата: {output_path}")
    doc.save(output_path)
    debug_log("WORD", "Файл успешно сохранен")


def main():

    debug_log("MAIN", "Запуск генератора по JSON")
    root = tk.Tk()
    root.withdraw()

    json_files = filedialog.askopenfilenames(
        title="Выберите JSON файлы с вопросами",
        filetypes=[("JSON", "*.json")]
    )

    if not json_files:
        debug_log("MAIN", "JSON файлы не выбраны, завершение работы")
        return

    debug_log("MAIN", f"Выбрано JSON файлов: {len(json_files)}")
    for index, json_path in enumerate(json_files, start=1):
        debug_log("MAIN", f"Обработка файла {index}/{len(json_files)}: {json_path}")

        data = load_questions_from_json(json_path)
        fallback_name = os.path.splitext(os.path.basename(json_path))[0]
        document_name = get_document_title(data, fallback_name)
        output_file_name = sanitize_filename(f"Тест_{document_name}") + ".docx"
        output_path = os.path.join(
            os.path.dirname(json_path),
            output_file_name
        )

        fill_word_template(
            TEMPLATE_PATH,
            output_path,
            document_name,
            data
        )

        print("Создан тест:", output_path, flush=True)
        debug_log("MAIN", f"Готово {index}/{len(json_files)}")

    debug_log("MAIN", "Все JSON файлы обработаны")


if __name__ == "__main__":
    main()
